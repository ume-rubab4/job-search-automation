"""
Build / refresh data/profile.json from an uploaded CV (PDF or DOCX).
The LLM extracts only what is written in the CV; nothing is invented.
"""
import io, json, shutil
from datetime import datetime
from .core import ROOT, PROFILE, ask_json

PROFILE_PATH = ROOT / "data" / "profile.json"

SCHEMA = {
    "name": "", "headline": "", "email": "", "phone": "", "location": "", "nationality": "",
    "needs_visa_sponsorship": True, "availability": "", "linkedin": "", "summary": "",
    "skills": [], "experience": [{"title": "", "company": "", "location": "", "start": "", "end": "", "bullets": []}],
    "education": [{"degree": "", "institution": "", "end": ""}],
    "certifications": [], "achievements": [], "languages": [],
}

SYSTEM = """You convert a CV into structured JSON for a job-matching agent.
Return ONLY JSON matching the given schema. Copy facts exactly as written; never invent, infer or embellish.
If a field is not present in the CV, use an empty string / empty list, except:
- needs_visa_sponsorship: keep the value given in CURRENT PROFILE unless the CV clearly states otherwise.
- nationality: keep CURRENT PROFILE value if the CV doesn't state one.
Keep bullet points as full sentences from the CV. Keep skills as short phrases."""


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    if name.endswith(".docx"):
        from docx import Document
        d = Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    return data.decode(errors="ignore")


def build_profile(filename: str, data: bytes) -> dict:
    text = extract_text(filename, data)
    if len(text.strip()) < 200:
        raise ValueError("Couldn't read enough text from that file. If it's a scanned PDF, export it as text first.")
    keep = {k: PROFILE.get(k) for k in ("needs_visa_sponsorship", "nationality", "linkedin", "availability")}
    user = (f"SCHEMA:\n{json.dumps(SCHEMA)}\n\nCURRENT PROFILE (for the fields mentioned in the rules):\n{json.dumps(keep)}\n\n"
            f"CV TEXT:\n{text[:15000]}")
    prof = ask_json(SYSTEM, user, max_tokens=4000)
    for k, v in keep.items():
        if not prof.get(k) and v:
            prof[k] = v
    return prof


def save_profile(prof: dict) -> str:
    backup = ROOT / "data" / f"profile_backup_{datetime.now():%Y%m%d_%H%M%S}.json"
    if PROFILE_PATH.exists():
        shutil.copy(PROFILE_PATH, backup)
    PROFILE_PATH.write_text(json.dumps(prof, indent=2, ensure_ascii=False))
    return str(backup)