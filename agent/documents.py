import json, re
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4, LETTER
from reportlab.lib.units import mm
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem, HRFlowable
from .core import PROFILE, OUTPUT_DIR, ask_json, ask_text
from .country_rules import for_country

TAILOR_SYSTEM = """You tailor a candidate's CV to one job posting.
HARD RULES:
- Never invent skills, tools, employers, dates, degrees, certifications, or achievements.
- You may reorder, rephrase, shorten, and emphasise existing items and mirror the posting's wording.
- Respect the country conventions given. Keep within the page limit by trimming bullets.
Return ONLY JSON with this shape:
{"headline": str, "summary": str,
 "skills": [str], "experience": [{"title":str,"company":str,"location":str,"start":str,"end":str,"bullets":[str]}],
 "education": [..same as input..], "certifications": [str], "achievements": [str], "languages": [str],
 "work_authorisation": str}"""


def tailor(job: dict) -> dict:
    rules = for_country(job["country"])
    user = (f"MASTER PROFILE (source of truth):\n{json.dumps(PROFILE, ensure_ascii=False)}\n\n"
            f"COUNTRY CONVENTIONS:\n{json.dumps(rules)}\n\n"
            f"JOB:\n{job['title']} at {job['company']} ({job.get('city')}, {job['country']})\n"
            f"{job['description'][:6000]}")
    return ask_json(TAILOR_SYSTEM, user, max_tokens=3000)


# ---------------------------------------------------------------- shared content

def _sections(job: dict, t: dict) -> dict:
    rules = for_country(job["country"])
    contact = [PROFILE["location"], PROFILE["phone"], PROFILE["email"]]
    if PROFILE.get("linkedin"):
        contact.append(PROFILE["linkedin"])
    personal = []
    if rules["nationality"]:
        personal.append(f"Nationality: {PROFILE['nationality']}")
    if rules["dob"]:
        personal.append("Date of birth: [ADD]")
    if rules["marital"]:
        personal.append("Marital status: [ADD]")
    if rules["photo"]:
        personal.append("[Insert professional photo top-right before sending]")
    return {
        "rules": rules, "contact": " | ".join(contact), "personal": " | ".join(personal),
        "auth": f"Work authorisation: {t.get('work_authorisation', 'Requires employment visa sponsorship')}",
        "headline": t.get("headline", PROFILE["headline"]),
        "signature": f"Lahore, [DATE]        {PROFILE['name']}" if rules["country"] in ("DE", "AT", "CH") else "",
    }


def _slug(job: dict) -> str:
    s = f"{job['company']}_{job['title']}_{job['country']}"
    return re.sub(r"[^A-Za-z0-9]+", "_", s)[:80]


# ---------------------------------------------------------------- DOCX

def _h(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text.upper()); r.bold = True; r.font.size = Pt(11.5)


def _build_docx(job: dict, t: dict, path):
    s = _sections(job, t)
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    name = doc.add_paragraph(); name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = name.add_run(PROFILE["name"].upper()); r.bold = True; r.font.size = Pt(16)
    hl = doc.add_paragraph(s["headline"]); hl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = doc.add_paragraph(s["contact"]); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if s["personal"]:
        doc.add_paragraph(s["personal"])
    doc.add_paragraph(s["auth"])
    _h(doc, s["rules"]["summary_heading"]); doc.add_paragraph(t["summary"])
    _h(doc, "Core skills")
    for x in t["skills"]:
        doc.add_paragraph(x, style="List Bullet")
    _h(doc, "Professional experience")
    for e in t["experience"]:
        p = doc.add_paragraph()
        p.add_run(f"{e['title']} - {e['company']}").bold = True
        p.add_run(f"  ({e.get('location','')})  {e['start']} - {e['end']}")
        for b in e["bullets"]:
            doc.add_paragraph(b, style="List Bullet")
    _h(doc, "Education")
    for ed in t["education"]:
        doc.add_paragraph(f"{ed['degree']}, {ed['institution']} ({ed.get('end','')})")
    _h(doc, "Certifications")
    for cert in t["certifications"]:
        doc.add_paragraph(cert, style="List Bullet")
    if t.get("achievements"):
        _h(doc, "Key achievements")
        for a in t["achievements"]:
            doc.add_paragraph(a, style="List Bullet")
    _h(doc, "Languages"); doc.add_paragraph(" | ".join(t["languages"]))
    if s["rules"]["references"] == "on request":
        doc.add_paragraph("References available on request")
    if s["signature"]:
        doc.add_paragraph("\n" + s["signature"])
    doc.save(path)


# ---------------------------------------------------------------- PDF

def _styles():
    base = ParagraphStyle("base", fontName="Helvetica", fontSize=10, leading=13)
    return {
        "base": base,
        "name": ParagraphStyle("name", parent=base, fontName="Helvetica-Bold", fontSize=17, leading=20, alignment=TA_CENTER, spaceAfter=2),
        "headline": ParagraphStyle("headline", parent=base, fontSize=10.5, alignment=TA_CENTER, textColor="#333333"),
        "contact": ParagraphStyle("contact", parent=base, fontSize=9, alignment=TA_CENTER, textColor="#444444", spaceAfter=6),
        "h": ParagraphStyle("h", parent=base, fontName="Helvetica-Bold", fontSize=10.5, leading=13, spaceBefore=9, spaceAfter=3, textColor="#1F2A44"),
        "role": ParagraphStyle("role", parent=base, fontName="Helvetica-Bold", spaceBefore=4),
        "meta": ParagraphStyle("meta", parent=base, fontSize=9, textColor="#555555"),
    }


def _esc(x: str) -> str:
    return (x or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _bullets(items, st):
    return ListFlowable([ListItem(Paragraph(_esc(i), st["base"]), leftIndent=10) for i in items],
                        bulletType="bullet", bulletFontSize=7, leftIndent=12, spaceBefore=1)


def _build_pdf(job: dict, t: dict, path):
    s = _sections(job, t); st = _styles()
    size = LETTER if s["rules"]["country"] == "US" else A4
    doc = SimpleDocTemplate(str(path), pagesize=size, leftMargin=18*mm, rightMargin=18*mm, topMargin=16*mm, bottomMargin=16*mm,
                            title=f"{PROFILE['name']} - CV", author=PROFILE["name"])
    f = [Paragraph(_esc(PROFILE["name"].upper()), st["name"]),
         Paragraph(_esc(s["headline"]), st["headline"]),
         Paragraph(_esc(s["contact"]), st["contact"])]
    if s["personal"]:
        f.append(Paragraph(_esc(s["personal"]), st["meta"]))
    f.append(Paragraph(_esc(s["auth"]), st["meta"]))
    f.append(HRFlowable(width="100%", thickness=0.6, color="#1F2A44", spaceBefore=6, spaceAfter=2))

    def h(txt): f.append(Paragraph(_esc(txt.upper()), st["h"]))

    h(s["rules"]["summary_heading"]); f.append(Paragraph(_esc(t["summary"]), st["base"]))
    h("Core skills"); f.append(_bullets(t["skills"], st))
    h("Professional experience")
    for e in t["experience"]:
        f.append(Paragraph(f"{_esc(e['title'])} — {_esc(e['company'])}", st["role"]))
        f.append(Paragraph(_esc(f"{e.get('location','')}  |  {e['start']} – {e['end']}"), st["meta"]))
        f.append(_bullets(e["bullets"], st))
    h("Education")
    for ed in t["education"]:
        f.append(Paragraph(_esc(f"{ed['degree']}, {ed['institution']} ({ed.get('end','')})"), st["base"]))
    h("Certifications"); f.append(_bullets(t["certifications"], st))
    if t.get("achievements"):
        h("Key achievements"); f.append(_bullets(t["achievements"], st))
    h("Languages"); f.append(Paragraph(_esc(" | ".join(t["languages"])), st["base"]))
    if s["rules"]["references"] == "on request":
        f += [Spacer(1, 6), Paragraph("References available on request", st["meta"])]
    if s["signature"]:
        f += [Spacer(1, 14), Paragraph(_esc(s["signature"]), st["base"])]
    doc.build(f)


def build_cv_docx(job: dict, t: dict) -> str:
    """Builds both formats; returns the PDF path (used for the digest and auto-apply). DOCX sits next to it."""
    base = OUTPUT_DIR / f"CV_{_slug(job)}"
    _build_docx(job, t, f"{base}.docx")
    _build_pdf(job, t, f"{base}.pdf")
    return f"{base}.pdf"


# ---------------------------------------------------------------- cover letter

COVER_SYSTEM = """Write a cover letter for the candidate for this exact job. Use only facts from the profile.
Follow the country cover-letter convention given. 250-350 words. Address the 2-3 most important
requirements in the posting with concrete evidence. Plain text, no placeholders except [DATE]."""


def build_cover_letter(job: dict, t: dict) -> str:
    rules = for_country(job["country"])
    user = (f"PROFILE:\n{json.dumps(PROFILE, ensure_ascii=False)}\n\nTAILORED SUMMARY:\n{t['summary']}\n\n"
            f"CONVENTION: {rules['cover_letter']}\nSPELLING: {rules['spelling']}\n\n"
            f"JOB: {job['title']} at {job['company']} ({job.get('city')}, {job['country']})\n{job['description'][:5000]}")
    text = ask_text(COVER_SYSTEM, user)
    base = OUTPUT_DIR / f"CoverLetter_{_slug(job)}"
    doc = Document(); doc.styles["Normal"].font.name = "Calibri"; doc.styles["Normal"].font.size = Pt(11)
    for para in text.split("\n\n"):
        doc.add_paragraph(para.strip())
    doc.save(f"{base}.docx")
    st = _styles()
    pdf = SimpleDocTemplate(f"{base}.pdf", pagesize=LETTER if job["country"] == "US" else A4,
                            leftMargin=20*mm, rightMargin=20*mm, topMargin=20*mm, bottomMargin=20*mm)
    body = ParagraphStyle("body", parent=st["base"], fontSize=10.5, leading=15, spaceAfter=8)
    pdf.build([Paragraph(_esc(p.strip()), body) for p in text.split("\n\n") if p.strip()])
    return f"{base}.pdf"